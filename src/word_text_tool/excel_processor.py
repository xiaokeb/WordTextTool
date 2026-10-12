"""Apply numbering through Office; verify independently before publication."""
import shutil
import tempfile
from pathlib import Path

from word_text_tool.file_service import prepare_output_directory
from word_text_tool.models import FileResult, FileStatus
from word_text_tool.numbering_service import NumberPreview, NumberSettings, preview_number
from word_text_tool.office_backend import run_office
from word_text_tool.office_validation import inspect_package
from word_text_tool.safe_output import OperationError, error_message, publish, require_unchanged


def number_workbook(book, number, fixed, fixed_text):
    for index in range(1, book.Worksheets.Count + 1):
        sheet = book.Worksheets.Item(index)
        if fixed:
            # Office shifts an existing print area down on insertion. Preserve
            # its column bounds and include the newly inserted heading when the
            # original area started at row 1; other areas merely move with data.
            print_areas = []
            if sheet.PageSetup.PrintArea:
                areas = sheet.Range(sheet.PageSetup.PrintArea).Areas
                for area_index in range(1, areas.Count + 1):
                    area = areas.Item(area_index)
                    print_areas.append((area.Row, area.Column, area.Rows.Count, area.Columns.Count))
            sheet.Rows.Item(1).Insert(Shift=-4121)
            if print_areas:
                addresses = []
                for top, left, rows, columns in print_areas:
                    first = 1 if top == 1 else top + 1
                    area = sheet.Range(sheet.Cells.Item(first, left), sheet.Cells.Item(top + rows, left + columns - 1))
                    addresses.append(area.Address)
                sheet.PageSetup.PrintArea = ",".join(addresses)
        values = [(1, fixed_text), (2, number)] if fixed else [(1, number)]
        for row, text in values:
            target = sheet.Range(f"E{row}:F{row}")
            target.Merge()
            target.NumberFormat = "@"
            target.HorizontalAlignment = -4108
            target.VerticalAlignment = -4108
            target.Font.Bold = True
            target.Font.Name = "Times New Roman"
            sheet.Cells.Item(row, 5).Value2 = text


def verify_numbered(path: Path, preview: NumberPreview, settings: NumberSettings):
    from openpyxl import load_workbook
    inspect_package(path)
    book = load_workbook(path)
    try:
        if tuple(book.sheetnames) != preview.sheets:
            raise OperationError("编号结果工作表不完整")
        values = [(1, settings.fixed_text), (2, preview.number)] if settings.fixed else [(1, preview.number)]
        for sheet in book:
            for row, value in values:
                cell = sheet.cell(row, 5)
                if (cell.value != value or cell.data_type == "f" or f"E{row}:F{row}" not in sheet.merged_cells
                        or cell.font.name != "Times New Roman" or not cell.font.bold
                        or cell.alignment.horizontal != "center" or cell.alignment.vertical != "center"):
                    raise OperationError(f"工作表“{sheet.title}”编号或格式验证失败")
    finally:
        book.close()


def write_number(preview: NumberPreview, settings: NumberSettings, output_dir=None, runner=run_office):
    source = preview.source
    require_unchanged(source, preview.fingerprint)
    preview_number(source, settings)  # Repeat safety preflight immediately before opening Office.
    if not settings.fixed:
        try:
            verify_numbered(source, preview, settings)
        except OperationError:
            pass
        else:
            return source, {"engine": settings.engine, "identity": "编号与格式无变化", "unchanged": True}
    directory = prepare_output_directory(source, output_dir)
    with tempfile.TemporaryDirectory(prefix=".wordtext-number-", dir=directory) as temporary:
        root = Path(temporary)
        shutil.copyfile(source, root / "input.xlsx")
        require_unchanged(source, preview.fingerprint)
        result = runner({"action": "number", "number": preview.number, "fixed": settings.fixed,
                         "fixed_text": settings.fixed_text, "engine": settings.engine}, root)
        numbered = root / "numbered.xlsx"
        verify_numbered(numbered, preview, settings)
        require_unchanged(source, preview.fingerprint)
        destination = publish(numbered, directory, source.stem + " - 编号版", ".xlsx")
    return destination, result


def process_excel(preview, settings, output_dir, action, cancelled, runner=run_office):
    from word_text_tool.pdf_processor import export_pdf
    outputs = []
    steps = []
    try:
        if action not in ("number", "pdf", "convert"):
            raise OperationError("未知的办公操作")
        if action != "pdf":
            settings.validate()
        current = preview.source
        engine = settings.engine
        if action in ("number", "convert"):
            current, result = write_number(preview, settings, output_dir, runner)
            engine = result["engine"]
            if result.get("unchanged"):
                steps.append("编号与格式无变化，未生成副本")
            else:
                outputs.append(current)
                steps.append(f"编号已保存（{result['identity']}）：{current}")
        if action in ("pdf", "convert"):
            if cancelled.is_set():
                return FileResult(preview.source, FileStatus.SKIPPED, "；".join(steps + ["PDF 已取消"]),
                                  outputs[0] if outputs else None, outputs=tuple(outputs))
            require_unchanged(preview.source, preview.fingerprint)
            pdf_directory = output_dir or (current.parent if outputs else None)
            pdf, result = export_pdf(current, pdf_directory, engine, runner,
                                     expected=preview.fingerprint if action == "pdf" else None)
            outputs.append(pdf)
            steps.append(f"PDF 已保存（{result['identity']}）：{pdf}")
        return FileResult(preview.source, FileStatus.SUCCESS if outputs else FileStatus.UNMATCHED,
                          "；".join(steps), outputs[-1] if outputs else None, outputs=tuple(outputs))
    except Exception as exc:
        return FileResult(preview.source, FileStatus.FAILED, "；".join(steps + [error_message(exc)]),
                          outputs[0] if outputs else None, outputs=tuple(outputs))
