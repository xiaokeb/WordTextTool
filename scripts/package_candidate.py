"""Attach reproducible provenance and an Excel acceptance sample to a candidate."""
import hashlib
import json
from pathlib import Path
import shutil
from candidate_files import file_digest, source_digest

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / "dist" / "V2.0-candidate"
exe = DEST / "Word 文本处理工具.exe"
assert exe.is_file(), "Build the candidate first"
digest = file_digest(exe)
receipt = json.loads((ROOT / "build" / "candidate-build.json").read_text(encoding="utf-8"))
if not receipt.get("ok") or receipt.get("exe_sha256") != digest or receipt.get("source_and_build_sha256") != source_digest():
    raise RuntimeError("Candidate does not match a successful current build; run build.bat")
manifest = {"version": "2.0-candidate", "exe": exe.name, "size": exe.stat().st_size,
            "sha256": digest, "source_and_build_sha256": receipt["source_and_build_sha256"],
            "built_at_utc": receipt["built_at_utc"],
            "release_status": "Windows 11/WPS validation required; Excel interface compatibility without real-engine acceptance; not published"}
(DEST / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
(DEST / "SHA256SUMS.txt").write_text(digest + "  " + exe.name + "\n", encoding="utf-8")
shutil.copyfile(ROOT / "docs" / "V2.0发布说明.md", DEST / "V2.0发布说明.md")
sample = Workbook()
sample.active.title = "第一张表"
sample.create_sheet("第二张表")
sample.create_sheet("最后一张表")
for sheet in sample:
    sheet["A1"] = "原始表头"
    sheet["A2"] = "中文内容验收"
    sheet["B3"] = 12
    sheet["B4"] = "=B3*2"
    sheet.merge_cells("E1:F1")
    sheet["E1"] = "旧编号"
    sheet.column_dimensions["A"].width = 35
    sheet.column_dimensions["E"].width = 22
    sheet.column_dimensions["F"].width = 22
    sheet.print_area = "A1:F5"
sample.save(DEST / "三表验收样本.xlsx")
(DEST / "README.md").write_text("""# Word 文本处理工具 V2.0 验收候选版

双击 `Word 文本处理工具.exe` 启动，无需安装 Python。
Word 文本替换与后缀管理可独立使用；记录编号和 PDF 导出需要本地安装可自动化的 WPS 表格或 Microsoft Excel。

本包为本次构建的验收候选版，具体 Windows 11/WPS 验证结果见随附发布说明。
Microsoft Excel 保留独立识别和兼容接口，但本机没有真实 Excel，未宣称完成真机验收。
使用附带的 `三表验收样本.xlsx`，按 [V2.0发布说明](V2.0发布说明.md) 的验收步骤操作。

固定模式：模板 `REC-yy`，起始 99，固定编号 `FIXED-编号`。先预览，再一键转换。
应在每张表顶端插入一行，新 E1 为固定编号，E2 为 REC-99；两处合并、居中、加粗、Times New Roman。
三页 PDF 均应显示两行编号，源工作簿不变。自动模式不插行，只写 E1:F1。

后缀管理会直接修改选中文件名；请先核对预览。所有页面均支持安全取消。
校验文件见 SHA256SUMS.txt，构建来源摘要见 manifest.json。
""", encoding="utf-8")
manifest["files"] = {path.name: {"sha256": file_digest(path), "size": path.stat().st_size}
                     for path in sorted(DEST.iterdir()) if path.is_file() and path.name not in ("manifest.json", "SHA256SUMS.txt")}
(DEST / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
(DEST / "SHA256SUMS.txt").write_text("".join(item["sha256"] + "  " + name + "\n" for name, item in manifest["files"].items()), encoding="utf-8")
print(json.dumps(manifest, ensure_ascii=True, indent=2))
