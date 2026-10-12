"""Real Office integration check. Writes only under build/office-verification."""
import argparse
import json
import sys
import uuid
import hashlib
from contextlib import nullcontext
from pathlib import Path
from threading import Event

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from word_text_tool.excel_processor import process_excel
from word_text_tool.numbering_service import NumberSettings, preview_number
from word_text_tool.office_backend import run_office, OfficeSession
from word_text_tool.safe_output import fingerprint


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--engine", choices=["wps", "excel", "auto"], required=True)
    parser.add_argument("--exe")
    parser.add_argument("--mode", choices=["fixed", "automatic"], default="fixed")
    parser.add_argument("--with-existing", action="store_true")
    args = parser.parse_args()
    root = ROOT / "build" / "office-verification" / uuid.uuid4().hex[:8]
    root.mkdir(parents=True)
    source = root / "多表编号验收.xlsx"
    workbook = Workbook()
    workbook.active.title = "第一张表"
    workbook.create_sheet("第二张表")
    workbook.create_sheet("最后一张表")
    for sheet in workbook:
        sheet["A1"] = "原始表头"
        sheet["A2"] = "中文内容与编号验收"
        sheet["B3"] = 12
        sheet["B4"] = "=B3*2"
        sheet.column_dimensions["A"].width = 35
        sheet.column_dimensions["E"].width = 22
        sheet.column_dimensions["F"].width = 22
        sheet.merge_cells("E1:F1")
        sheet["E1"] = "旧编号"
        sheet.print_area = "A1:F5"
    workbook.save(source)
    fixed = args.mode == "fixed"
    settings = NumberSettings("REC-yy", 99, fixed, "FIXED-编号", args.engine)
    preview = preview_number(source, settings)
    calls = []

    def runner(request, directory):
        request["diagnostics"] = True
        try:
            result = run_office(request, directory, timeout=45, executable=args.exe)
            calls.append(result)
            return result
        except Exception:
            result_file = directory / "result.json"
            if result_file.exists():
                (root / "diagnostic.json").write_bytes(result_file.read_bytes())
            raise

    guard = OfficeSession(args.engine, root / "existing-owner.json") if args.with_existing else nullcontext()
    with guard as existing:
        if existing:
            existing.open(source)
        result = process_excel(preview, settings, root, "convert", Event(), runner)
        if existing:
            assert existing.app.Workbooks.Count == 1
            assert existing.book.Worksheets.Item(1).Cells.Item(1, 5).Value2 == "旧编号"
    report = {"engine": args.engine, "mode": args.mode, "status": result.status.value, "message": result.message,
              "outputs": [str(p) for p in result.outputs], "calls": calls,
              "source_unchanged": fingerprint(source) == preview.fingerprint}
    report["existing_workbook_preserved"] = args.with_existing
    if args.exe:
        report["candidate_sha256"] = hashlib.sha256(Path(args.exe).read_bytes()).hexdigest().upper()
    if result.status.value == "success":
        book = load_workbook(result.outputs[0])
        assert len(book.worksheets) == 3
        for sheet in book:
            assert sheet["A2" if fixed else "A1"].value == "原始表头"
            assert sheet["B5" if fixed else "B4"].value == ("=B4*2" if fixed else "=B3*2")
            assert sheet["E2" if fixed else "E1"].value == "REC-99"
            if fixed:
                assert sheet["E1"].value == "FIXED-编号"
        book.close()
        pdf = PdfReader(result.outputs[1])
        assert len(pdf.pages) >= 3
        text = "\n".join(page.extract_text() for page in pdf.pages)
        assert text.count("REC-99") >= 3
        if fixed:
            assert text.count("FIXED-") >= 3, "Inserted fixed header missing from PDF print area"
        report["pdf_pages"] = len(pdf.pages)
        report["all_sheets_verified"] = True
    (root / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=True, indent=2))
    print("REPORT:", root)
    return 0 if result.status.value == "success" else 1


if __name__ == "__main__":
    raise SystemExit(main())
