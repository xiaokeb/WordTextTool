from pathlib import Path
from threading import Event

import pytest
from docx import Document
from PyQt6 import sip
from PyQt6.QtCore import QThread, QTimer, pyqtSignal
from PyQt6.QtWidgets import QApplication, QMessageBox

from word_text_tool.app import create_application
from word_text_tool.models import BatchSummary, FileStatus
from word_text_tool.ui.main_window import MainWindow


class PausingAfterSummaryWorker(QThread):
    """Real QThread that exposes the gap between its two completion signals."""

    progress = pyqtSignal(int, int, str, int)
    batch_finished = pyqtSignal(object)

    def __init__(self, processor, files, rules, custom_output_dir, parent=None):
        super().__init__(parent)
        self.processor = processor
        self.files = files
        self.rules = rules
        self.custom_output_dir = custom_output_dir
        self.release = Event()
        self.cancel_requested = Event()
        self.native_finished = Event()
        self.finished.connect(self.native_finished.set)

    def request_cancel(self):
        self.cancel_requested.set()

    def run(self):
        self.progress.emit(1, 1, self.files[0].name, 100)
        self.batch_finished.emit(BatchSummary(()))
        self.release.wait(timeout=5)


def _seed_window(window, source):
    window.file_panel.add_paths((source,))
    window.rule_panel.before_input.setPlainText("旧值")
    window.rule_panel.after_input.setPlainText("新值")
    window.rule_panel.add_button.click()


def _stop_worker(window, worker):
    if not sip.isdeleted(worker):
        worker.release.set()
        worker.request_cancel()
        worker.wait(5000)
    if window.worker is None:
        window.close()


@pytest.mark.parametrize("width,height", [(820, 700), (1000, 900), (1400, 1000)])
def test_responsive_sizes_keep_actions_visible_and_rule_columns_balanced(
    qtbot, width, height
):
    window = MainWindow()
    qtbot.addWidget(window)
    window.show()
    window.resize(width, height)
    qtbot.wait(30)

    assert (window.width(), window.height()) == (width, height)
    for button in (window.start_button, window.cancel_button, window.clear_button):
        assert button.isVisibleTo(window)
        assert button.width() >= 110
        assert button.height() == 36
    header = window.rule_panel.table.horizontalHeader()
    assert abs(header.sectionSize(0) - header.sectionSize(1)) <= 2


def test_start_reports_each_empty_precondition_without_creating_worker(
    qtbot, tmp_path, monkeypatch
):
    source = tmp_path / "合同.docx"
    source.write_bytes(b"x")
    window = MainWindow()
    qtbot.addWidget(window)
    warnings = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: warnings.append(args))

    window.start_button.click()
    assert window.worker is None
    assert "DOCX" in warnings[-1][2]

    window.file_panel.add_paths((source,))
    window.start_button.click()
    assert window.worker is None
    assert "替换规则" in warnings[-1][2]
    assert len(warnings) == 2


def test_clear_requires_confirmation_and_empty_clear_does_not_prompt(
    qtbot, tmp_path, monkeypatch
):
    source = tmp_path / "合同.docx"
    source.write_bytes(b"x")
    window = MainWindow()
    qtbot.addWidget(window)
    _seed_window(window, source)
    window.custom_output_dir = tmp_path
    window.output_path_line_edit.setText(str(tmp_path))
    window.progress_bar.setValue(67)
    answers = [QMessageBox.StandardButton.No, QMessageBox.StandardButton.Yes]
    questions = []

    def answer(*args):
        questions.append(args)
        return answers.pop(0)

    monkeypatch.setattr(QMessageBox, "question", answer)

    window.clear_button.click()
    assert window.file_panel.files() == (source.resolve(),)
    assert len(window.rule_panel.rules()) == 1
    assert window.custom_output_dir == tmp_path

    window.clear_button.click()
    assert window.file_panel.files() == ()
    assert window.rule_panel.rules() == ()
    assert window.custom_output_dir is None
    assert window.progress_bar.value() == 0
    assert window.status_label.text() == "准备就绪"
    assert len(questions) == 2

    window.clear_button.click()
    assert len(questions) == 2


def test_batch_summary_does_not_unlock_or_destroy_worker_before_real_finished(
    qtbot, tmp_path, monkeypatch
):
    source = tmp_path / "合同.docx"
    source.write_bytes(b"x")
    window = MainWindow(worker_factory=PausingAfterSummaryWorker)
    qtbot.addWidget(window)
    shown = []
    monkeypatch.setattr(window, "_show_result_dialog", shown.append)
    _seed_window(window, source)

    window.start_button.click()
    worker = window.worker
    try:
        qtbot.waitUntil(lambda: window._worker_summary is not None, timeout=3000)

        assert worker.isRunning()
        assert window.worker is worker
        assert not sip.isdeleted(worker)
        assert not worker.native_finished.is_set()
        assert not window.start_button.isEnabled()
        assert window.cancel_button.isEnabled()
        assert shown == []
        assert worker.files == (source.resolve(),)
        assert isinstance(worker.rules, tuple)

        worker.release.set()
        qtbot.waitUntil(lambda: window.worker is None, timeout=3000)
        assert worker.native_finished.is_set()
        assert window.start_button.isEnabled()
        assert not window.cancel_button.isEnabled()
        assert len(shown) == 1
    finally:
        _stop_worker(window, worker)


def test_confirmed_close_waits_until_real_qthread_finished(qtbot, tmp_path, monkeypatch):
    source = tmp_path / "合同.docx"
    source.write_bytes(b"x")
    window = MainWindow(worker_factory=PausingAfterSummaryWorker)
    qtbot.addWidget(window)
    window.show()
    _seed_window(window, source)
    monkeypatch.setattr(
        QMessageBox, "question", lambda *args: QMessageBox.StandardButton.Yes
    )
    window.start_button.click()
    worker = window.worker
    try:
        qtbot.waitUntil(lambda: window._worker_summary is not None, timeout=3000)

        window.close()
        assert window.isVisible()
        assert worker.isRunning()
        assert worker.cancel_requested.is_set()

        worker.release.set()
        qtbot.waitUntil(lambda: not window.isVisible(), timeout=3000)
        assert worker.native_finished.is_set()
    finally:
        _stop_worker(window, worker)


def test_rejected_close_keeps_processing_and_later_finishes_normally(
    qtbot, tmp_path, monkeypatch
):
    source = tmp_path / "合同.docx"
    source.write_bytes(b"x")
    window = MainWindow(worker_factory=PausingAfterSummaryWorker)
    qtbot.addWidget(window)
    window.show()
    shown = []
    monkeypatch.setattr(window, "_show_result_dialog", shown.append)
    monkeypatch.setattr(
        QMessageBox, "question", lambda *args: QMessageBox.StandardButton.No
    )
    _seed_window(window, source)
    window.start_button.click()
    worker = window.worker
    try:
        qtbot.waitUntil(lambda: window._worker_summary is not None, timeout=3000)

        window.close()
        assert window.isVisible()
        assert worker.isRunning()
        assert not worker.cancel_requested.is_set()

        worker.release.set()
        qtbot.waitUntil(lambda: window.worker is None, timeout=3000)
        assert worker.native_finished.is_set()
        assert window.isVisible()
        assert len(shown) == 1
    finally:
        _stop_worker(window, worker)


@pytest.mark.parametrize("answer", [QMessageBox.StandardButton.Yes, QMessageBox.StandardButton.No])
def test_close_answer_after_worker_finalized_in_real_confirmation_loop(
    qtbot, tmp_path, answer
):
    source = tmp_path / "合同.docx"
    source.write_bytes(b"x")
    window = MainWindow(worker_factory=PausingAfterSummaryWorker)
    qtbot.addWidget(window)
    window.show()
    _seed_window(window, source)
    window.start_button.click()
    worker = window.worker
    answered_after_finalizing = []
    timer = QTimer(window)

    def complete_then_answer():
        dialogs = [item for item in window.findChildren(QMessageBox) if item.isVisible()]
        if not dialogs:
            return
        worker.release.set()
        if window.worker is None:
            answered_after_finalizing.append(True)
            timer.stop()
            dialogs[0].button(answer).click()

    timer.timeout.connect(complete_then_answer)
    try:
        qtbot.waitUntil(lambda: window._worker_summary is not None, timeout=3000)
        timer.start(10)
        window.close()
        assert answered_after_finalizing == [True]
        assert window.worker is None
        assert window.isVisible() == (answer == QMessageBox.StandardButton.No)
        assert not window._close_after_worker
    finally:
        timer.stop()
        _stop_worker(window, worker)


def test_real_main_window_processes_docx_without_changing_source_or_overwriting(
    qtbot, tmp_path, monkeypatch
):
    source = tmp_path / "合同.docx"
    document = Document()
    document.add_paragraph("合同中的旧值")
    document.save(source)
    source_bytes = source.read_bytes()
    output_dir = tmp_path / "修改版"
    output_dir.mkdir()
    existing = output_dir / "合同 - 修改版.docx"
    existing.write_bytes(b"existing output must stay unchanged")
    window = MainWindow()
    qtbot.addWidget(window)
    shown = []
    monkeypatch.setattr(window, "_show_result_dialog", shown.append)
    _seed_window(window, source)

    window.start_button.click()
    worker = window.worker
    try:
        qtbot.waitUntil(lambda: window.worker is None, timeout=5000)
    finally:
        if not sip.isdeleted(worker) and worker.isRunning():
            worker.request_cancel()
            worker.wait(5000)

    assert source.read_bytes() == source_bytes
    assert existing.read_bytes() == b"existing output must stay unchanged"
    assert len(shown) == 1
    result = shown[0].results[0]
    assert result.status is FileStatus.SUCCESS
    assert result.output == output_dir / "合同 - 修改版 (1).docx"
    assert result.output.exists()
    assert Document(result.output).paragraphs[0].text == "合同中的新值"
    assert window.progress_bar.value() == 100
    assert window.status_label.text() == "处理完成"


def test_create_application_sets_identity_and_font(qapp):
    application = create_application([])

    assert application is QApplication.instance()
    assert application.applicationName() == "Word 文本处理工具"
    assert application.organizationName() == "WordTextTool"
    assert application.font().family() == "Microsoft YaHei UI"
    assert application.font().pointSize() == 10
