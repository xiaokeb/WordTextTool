import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches
from docx.text.run import Run
from lxml import etree

from word_text_tool.document_processor import replace_document
from word_text_tool.models import ReplacementRule


RULES = (ReplacementRule("旧值", "新值"), ReplacementRule("新值", "再次替换"))


def add_hyperlink(paragraph, *texts):
    rid = paragraph.part.relate_to("https://example.com/old", RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    hyperlink.set(qn("w:history"), "1")
    paragraph._p.append(hyperlink)
    for text in texts:
        element = OxmlElement("w:r")
        hyperlink.append(element)
        run = Run(element, paragraph)
        run.text = text
        run.underline = True
    return hyperlink, rid


def add_field_marker(container, marker_type):
    run = OxmlElement("w:r")
    marker = OxmlElement("w:fldChar")
    marker.set(qn("w:fldCharType"), marker_type)
    run.append(marker)
    container.append(run)
    return run


def add_wrapped_field_marker(container, marker_type, wrapper_tag):
    wrapper = OxmlElement(wrapper_tag)
    marker_container = wrapper
    if wrapper_tag == "w:sdt":
        marker_container = OxmlElement("w:sdtContent")
        wrapper.append(marker_container)
    add_field_marker(marker_container, marker_type)
    container.append(wrapper)
    return wrapper


def add_xml_run(container, text):
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    container.append(run)
    return run


def text_of(element):
    return "".join(node.text or "" for node in element.xpath(".//w:t"))


def xml_of(element):
    return etree.tostring(element)


def test_body_tables_and_nested_tables_are_processed():
    document = Document()
    paragraph = document.add_paragraph("旧值正文")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "旧值表格"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "旧值嵌套"
    assert replace_document(document, RULES) == 3
    assert paragraph.text == "新值正文"
    assert table.cell(0, 0).paragraphs[0].text == "新值表格"
    assert nested.cell(0, 0).text == "新值嵌套"


@pytest.mark.parametrize("story_name", [
    "header", "footer", "first_page_header", "first_page_footer",
    "even_page_header", "even_page_footer",
])
def test_all_referenced_story_types_include_nested_tables_and_are_processed_once(story_name):
    document = Document()
    story = getattr(document.sections[0], story_name)
    story.paragraphs[0].text = "旧值"
    table = story.add_table(rows=1, cols=1, width=Inches(3))
    table.cell(0, 0).text = "旧值"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "旧值"
    document.add_section()
    assert replace_document(document, RULES) == 3
    assert story.paragraphs[0].text == "新值"
    assert table.cell(0, 0).paragraphs[0].text == "新值"
    assert nested.cell(0, 0).text == "新值"


def test_explicit_shared_part_references_are_deduplicated_by_part():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "旧值"
    rid = document.sections[0]._sectPr.headerReference_lst[0].rId
    second = document.add_section()
    reference = OxmlElement("w:headerReference")
    reference.set(qn("w:type"), "first")
    reference.set(qn("r:id"), rid)
    second._sectPr.insert(0, reference)
    assert replace_document(document, RULES) == 1
    assert second.first_page_header.paragraphs[0].text == "新值"


def test_absent_stories_are_not_created():
    document = Document()
    document.add_paragraph("旧值")
    relations_before = tuple(document.part.rels)
    assert replace_document(document, RULES) == 1
    assert tuple(document.part.rels) == relations_before
    assert not document._element.xpath(".//w:headerReference | .//w:footerReference")


@pytest.mark.parametrize("rows,cols", [(1, 3), (3, 1), (2, 2)])
def test_merged_cell_is_processed_once(rows, cols):
    document = Document()
    table = document.add_table(rows=rows, cols=cols)
    merged = table.cell(0, 0).merge(table.cell(rows - 1, cols - 1))
    merged.text = "旧值"
    assert replace_document(document, RULES) == 1
    assert merged.text == "新值"


def test_hyperlink_matches_across_its_runs_and_preserves_relationship_and_style():
    document = Document()
    paragraph = document.add_paragraph("前缀")
    hyperlink, rid = add_hyperlink(paragraph, "旧", "值后缀")
    styles = [run.rPr.xml for run in hyperlink.xpath("./w:r")]
    attributes = dict(hyperlink.attrib)
    assert replace_document(document, RULES) == 1
    assert text_of(hyperlink) == "新值后缀"
    assert dict(hyperlink.attrib) == attributes
    assert paragraph.part.rels[rid].target_ref == "https://example.com/old"
    assert [run.rPr.xml for run in hyperlink.xpath("./w:r")] == styles


@pytest.mark.parametrize("layout", ["normal-link", "link-normal", "link-link"])
def test_matches_never_cross_hyperlink_boundaries(layout):
    document = Document()
    paragraph = document.add_paragraph()
    if layout == "normal-link":
        paragraph.add_run("旧")
        add_hyperlink(paragraph, "值")
    elif layout == "link-normal":
        add_hyperlink(paragraph, "旧")
        paragraph.add_run("值")
    else:
        add_hyperlink(paragraph, "旧")
        add_hyperlink(paragraph, "值")
    before = paragraph._p.xml
    assert replace_document(document, RULES) == 0
    assert paragraph._p.xml == before


def test_simple_field_xml_is_unchanged_and_separates_normal_text():
    document = Document()
    paragraph = document.add_paragraph("旧")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " DATE 旧值 ")
    add_xml_run(field, "旧值")
    paragraph._p.append(field)
    paragraph.add_run("值")
    before = paragraph._p.xml
    assert replace_document(document, RULES) == 0
    assert paragraph._p.xml == before


@pytest.mark.parametrize("inside_hyperlink", [False, True])
def test_nested_complex_field_code_and_results_are_preserved(inside_hyperlink):
    document = Document()
    paragraph = document.add_paragraph()
    container = add_hyperlink(paragraph)[0] if inside_hyperlink else paragraph._p
    add_field_marker(container, "begin")
    instruction = OxmlElement("w:r")
    code = OxmlElement("w:instrText")
    code.text = " DATE 旧值 "
    instruction.append(code)
    container.append(instruction)
    add_field_marker(container, "separate")
    add_xml_run(container, "旧值")
    add_field_marker(container, "begin")
    add_field_marker(container, "separate")
    add_xml_run(container, "旧值")
    add_field_marker(container, "end")
    add_xml_run(container, "旧值")
    add_field_marker(container, "end")
    preserved = [node.xml for node in container]
    editable = add_xml_run(container, "旧值")
    assert replace_document(document, RULES) == 1
    assert [node.xml for node in list(container)[:-1]] == preserved
    assert text_of(editable) == "新值"


def test_complex_field_continues_across_paragraphs_and_table_in_story_order():
    document = Document()
    first = document.add_paragraph()
    add_field_marker(first._p, "begin")
    add_field_marker(first._p, "separate")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "旧值"
    last = document.add_paragraph("旧值")
    add_field_marker(last._p, "end")
    editable = last.add_run("旧值")
    assert replace_document(document, RULES) == 1
    assert table.cell(0, 0).text == "旧值"
    assert last.runs[0].text == "旧值"
    assert editable.text == "新值"


def test_complex_field_enclosing_hyperlink_keeps_link_result_unchanged():
    document = Document()
    paragraph = document.add_paragraph()
    add_field_marker(paragraph._p, "begin")
    add_field_marker(paragraph._p, "separate")
    hyperlink, _ = add_hyperlink(paragraph, "旧值")
    add_field_marker(paragraph._p, "end")
    before = hyperlink.xml
    paragraph.add_run("旧值")
    assert replace_document(document, RULES) == 1
    assert hyperlink.xml == before


@pytest.mark.parametrize(
    "begin_wrapper,end_wrapper",
    [("w:ins", None), (None, "w:smartTag")],
)
def test_wrapped_field_boundary_protects_result_and_releases_following_text(
    begin_wrapper, end_wrapper
):
    document = Document()
    paragraph = document.add_paragraph()
    if begin_wrapper:
        begin = add_wrapped_field_marker(paragraph._p, "begin", begin_wrapper)
    else:
        begin = add_field_marker(paragraph._p, "begin")
    add_field_marker(paragraph._p, "separate")
    result = add_xml_run(paragraph._p, "旧值")
    if end_wrapper:
        end = add_wrapped_field_marker(paragraph._p, "end", end_wrapper)
    else:
        end = add_field_marker(paragraph._p, "end")
    editable = add_xml_run(paragraph._p, "旧值")
    boundary_xml = (xml_of(begin), xml_of(end))

    assert replace_document(document, RULES) == 1
    assert text_of(result) == "旧值"
    assert text_of(editable) == "新值"
    assert (xml_of(begin), xml_of(end)) == boundary_xml


def test_nested_fields_track_markers_inside_revision_smart_tag_and_content_control():
    document = Document()
    paragraph = document.add_paragraph()
    wrappers = [add_wrapped_field_marker(paragraph._p, "begin", "w:ins")]
    add_field_marker(paragraph._p, "separate")
    protected = [add_xml_run(paragraph._p, "旧值")]
    wrappers.append(add_wrapped_field_marker(paragraph._p, "begin", "w:smartTag"))
    add_field_marker(paragraph._p, "separate")
    protected.append(add_xml_run(paragraph._p, "旧值"))
    wrappers.append(add_wrapped_field_marker(paragraph._p, "end", "w:sdt"))
    protected.append(add_xml_run(paragraph._p, "旧值"))
    wrappers.append(add_wrapped_field_marker(paragraph._p, "end", "w:smartTag"))
    editable = add_xml_run(paragraph._p, "旧值")
    wrapper_xml = [xml_of(wrapper) for wrapper in wrappers]

    assert replace_document(document, RULES) == 1
    assert [text_of(run) for run in protected] == ["旧值"] * 3
    assert text_of(editable) == "新值"
    assert [xml_of(wrapper) for wrapper in wrappers] == wrapper_xml


def test_textbox_field_marker_does_not_change_surrounding_story_field_state():
    document = Document()
    paragraph = document.add_paragraph()
    drawing_run = OxmlElement("w:r")
    drawing = OxmlElement("w:drawing")
    textbox = OxmlElement("w:txbxContent")
    textbox_paragraph = OxmlElement("w:p")
    add_field_marker(textbox_paragraph, "begin")
    textbox.append(textbox_paragraph)
    drawing.append(textbox)
    drawing_run.append(drawing)
    paragraph._p.append(drawing_run)
    editable = add_xml_run(paragraph._p, "旧值")

    assert replace_document(document, RULES) == 1
    assert text_of(editable) == "新值"


@pytest.mark.parametrize("tag", ["w:drawing", "w:object", "w:pict", "w:footnoteReference"])
@pytest.mark.parametrize("inside_hyperlink", [False, True])
def test_unsafe_run_xml_is_preserved_and_separates_matching_regions(tag, inside_hyperlink):
    document = Document()
    paragraph = document.add_paragraph()
    container = add_hyperlink(paragraph)[0] if inside_hyperlink else paragraph._p
    add_xml_run(container, "旧")
    unsafe = add_xml_run(container, "旧值")
    unsafe.append(OxmlElement(tag))
    add_xml_run(container, "值")
    before = container.xml
    assert replace_document(document, RULES) == 0
    assert container.xml == before


@pytest.mark.parametrize("break_type", [WD_BREAK.PAGE, WD_BREAK.COLUMN, WD_BREAK.LINE_CLEAR_ALL])
def test_nonordinary_breaks_are_preserved_and_form_boundaries(break_type):
    document = Document()
    paragraph = document.add_paragraph("旧")
    unsafe = paragraph.add_run("旧值")
    unsafe.add_break(break_type)
    paragraph.add_run("值")
    before = paragraph._p.xml
    assert replace_document(document, RULES) == 0
    assert paragraph._p.xml == before


def test_normal_runs_match_across_tabs_empty_runs_and_soft_breaks_with_format():
    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("旧")
    first.bold = True
    paragraph.add_run("")
    paragraph.add_run("\t值\n")
    suffix = paragraph.add_run("末尾")
    suffix.italic = True
    properties = first._r.rPr.xml
    assert replace_document(document, (ReplacementRule("旧\t值\n末", "新\n"),)) == 1
    assert paragraph.text == "新\n尾"
    assert first._r.rPr.xml == properties
    assert suffix.text == "尾"
    assert suffix.italic is True


def test_no_match_preserves_all_document_xml_and_paragraph_boundary():
    document = Document()
    document.add_paragraph("旧")
    document.add_paragraph("值")
    before = document._element.xml
    assert replace_document(document, RULES) == 0
    assert document._element.xml == before
