from pathlib import Path
from threading import Event
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from pypdf import PdfWriter

from word_text_tool.numbering_service import NumberSettings, preview_number
from word_text_tool.excel_processor import process_excel, verify_numbered, number_workbook
from word_text_tool.models import FileStatus
from word_text_tool.office_validation import inspect_package, xml
from word_text_tool.safe_output import OperationError, fingerprint


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "模板.xlsx"
    book = Workbook()
    book.active.title = "第一表"
    book.create_sheet("最后表")
    for sheet in book:
        sheet["A1"] = "原始标题"
        sheet["A2"] = "保留内容"
        sheet["E1"] = "旧编号"
        sheet.merge_cells("E1:F1")
    book.save(path)
    return path


def simulated_office(request, root):
    """A filesystem adapter double. Real COM is separately verified by script."""
    if request["action"] == "number":
        book = load_workbook(root / "input.xlsx")
        for sheet in book:
            if request["fixed"]:
                sheet.unmerge_cells("E1:F1")
                sheet.insert_rows(1)
                values = [(1, request["fixed_text"]), (2, request["number"])]
            else:
                values = [(1, request["number"])]
            for row, value in values:
                cell = sheet.cell(row, 5)
                cell.value = value
                cell.data_type = "s"
                cell.font = Font(name="Times New Roman", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        book.save(root / "numbered.xlsx")
        book.close()
    else:
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        writer.write(root / "export.pdf")
    return {"engine": "wps", "identity": "测试适配器"}


def test_preview_uses_every_sheet_and_sequence_over_99(workbook):
    settings = NumberSettings("项目-yy", 99, True, "固定")
    preview = preview_number(workbook, settings, 1)
    assert preview.number == "项目-100"
    assert preview.sheets == ("第一表", "最后表")
    assert all("!E2" in message for message in preview.overwrites)
    assert fingerprint(workbook) == preview.fingerprint


@pytest.mark.parametrize("settings", [NumberSettings("无占位"), NumberSettings("yy-yy"), NumberSettings(start=0),
                                    NumberSettings(fixed=True), NumberSettings(engine="unknown")])
def test_invalid_number_settings(settings):
    with pytest.raises(OperationError):
        settings.validate()


@pytest.mark.parametrize("damage, reason", [("hidden", "不可见"), ("protected", "保护"),
                                          ("formula", "公式"), ("merge", "合并"), ("nonempty_f", "F1")])
def test_preview_rejects_unsafe_target_in_last_sheet(workbook, damage, reason):
    book = load_workbook(workbook)
    last = book.worksheets[-1]
    if damage == "hidden":
        last.sheet_state = "hidden"
    elif damage == "protected":
        last.protection.sheet = True
    elif damage == "formula":
        last["E1"] = "=1+1"
    else:
        last.unmerge_cells("E1:F1")
        if damage == "merge":
            last.merge_cells("D1:F2")
        else:
            last["F1"] = "不可丢弃"
    book.save(workbook)
    with pytest.raises(OperationError, match=reason):
        preview_number(workbook, NumberSettings())


@pytest.mark.parametrize("part", ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml", "xl/connections.xml", "xl/embeddings/oleObject1.bin"])
def test_active_content_rejected_before_office(workbook, part):
    with ZipFile(workbook, "a") as archive:
        archive.writestr(part, b"dangerous")
    with pytest.raises(OperationError):
        inspect_package(workbook)


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_xml_entities_rejected(encoding):
    data = '<?xml version="1.0"?><!DOCTYPE a [<!ENTITY x "payload">]><a>&x;</a>'.encode(encoding)
    with pytest.raises(OperationError):
        xml(data)


def test_suspicious_compression_rejected(tmp_path):
    path = tmp_path / "bomb.xlsx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"x" * 2 * 1024 * 1024)
    with pytest.raises(OperationError, match="压缩"):
        inspect_package(path)


def test_fixed_pipeline_preserves_source_and_routes_numbered_copy_to_pdf(workbook, tmp_path):
    settings = NumberSettings("项目-yy", 99, True, "固定编号")
    preview = preview_number(workbook, settings)
    calls = []

    def runner(request, root):
        calls.append(request["action"])
        if request["action"] == "pdf":
            book = load_workbook(root / "input.xlsx")
            for sheet in book:
                assert sheet["E1"].value == "固定编号"
                assert sheet["E2"].value == "项目-99"
                assert sheet["A2"].value == "原始标题"
                assert sheet["A3"].value == "保留内容"
            book.close()
        return simulated_office(request, root)

    result = process_excel(preview, settings, tmp_path, "convert", Event(), runner)
    assert result.status == FileStatus.SUCCESS
    assert calls == ["number", "pdf"]
    assert len(result.outputs) == 2
    assert all(p.is_file() for p in result.outputs)
    assert fingerprint(workbook) == preview.fingerprint
    verify_numbered(result.outputs[0], preview, settings)


def test_pdf_failure_keeps_numbered_copy(workbook, tmp_path):
    settings = NumberSettings()
    preview = preview_number(workbook, settings)
    def runner(request, root):
        if request["action"] == "pdf":
            raise OperationError("导出失败")
        return simulated_office(request, root)
    result = process_excel(preview, settings, tmp_path, "convert", Event(), runner)
    assert result.status == FileStatus.FAILED
    assert len(result.outputs) == 1
    assert result.outputs[0].exists()
    assert "编号已保存" in result.message and "导出失败" in result.message


def test_cancel_between_stages_keeps_numbered_copy(workbook, tmp_path):
    settings = NumberSettings()
    preview = preview_number(workbook, settings)
    cancel = Event()
    def runner(request, root):
        result = simulated_office(request, root)
        cancel.set()
        return result
    result = process_excel(preview, settings, tmp_path, "convert", cancel, runner)
    assert result.status == FileStatus.SKIPPED
    assert len(result.outputs) == 1 and "PDF 已取消" in result.message


def test_source_change_blocks_office(workbook, tmp_path):
    settings = NumberSettings()
    preview = preview_number(workbook, settings)
    workbook.write_bytes(b"changed")
    def unexpected(*args):
        pytest.fail("changed input reached Office")
    result = process_excel(preview, settings, tmp_path, "number", Event(), unexpected)
    assert result.status == FileStatus.FAILED
    assert "变化" in result.message


def test_invalid_saved_result_never_published(workbook, tmp_path):
    preview = preview_number(workbook, NumberSettings())
    def runner(request, root):
        (root / "numbered.xlsx").write_bytes(b"broken")
        return {"engine": "wps", "identity": "测试"}
    result = process_excel(preview, NumberSettings(), tmp_path, "number", Event(), runner)
    assert result.status == FileStatus.FAILED
    assert not list(tmp_path.glob("*编号版*"))
    assert not list(tmp_path.glob(".wordtext-*"))


def test_repeated_output_never_overwrites(workbook, tmp_path):
    preview = preview_number(workbook, NumberSettings())
    first = process_excel(preview, NumberSettings(), tmp_path, "number", Event(), simulated_office)
    snapshot = fingerprint(first.output)
    second = process_excel(preview, NumberSettings(), tmp_path, "number", Event(), simulated_office)
    assert first.status == second.status == FileStatus.SUCCESS
    assert first.output != second.output
    assert fingerprint(first.output) == snapshot


def test_formula_like_number_is_stored_as_text(workbook, tmp_path):
    settings = NumberSettings("=yy")
    preview = preview_number(workbook, settings)
    result = process_excel(preview, settings, tmp_path, "number", Event(), simulated_office)
    assert result.status == FileStatus.SUCCESS
    verify_numbered(result.output, preview, settings)


def test_one_click_default_directory_does_not_nest_modified_folder(workbook):
    preview = preview_number(workbook, NumberSettings())
    result = process_excel(preview, NumberSettings(), None, "convert", Event(), simulated_office)
    assert result.status == FileStatus.SUCCESS
    assert {p.parent for p in result.outputs} == {workbook.parent / "修改版"}


def test_pdf_only_does_not_require_valid_number_template(workbook, tmp_path):
    settings = NumberSettings(template="", fixed=True)
    preview = preview_number(workbook, settings, pdf_only=True)
    result = process_excel(preview, settings, tmp_path, "pdf", Event(), simulated_office)
    assert result.status == FileStatus.SUCCESS


def test_already_numbered_document_does_not_generate_another_copy(workbook, tmp_path):
    settings = NumberSettings()
    first = process_excel(preview_number(workbook, settings), settings, tmp_path, "number", Event(), simulated_office)
    second = process_excel(preview_number(first.output, settings), settings, tmp_path, "number", Event(), simulated_office)
    assert second.status == FileStatus.UNMATCHED
    assert not second.outputs


@pytest.mark.parametrize("formula", ['=WEBSERVICE("https://example.com")', '=RTD("external",,"data")', '=_xlfn.IMAGE("https://example.com")'])
def test_external_formula_outside_number_cells_rejected(workbook, formula):
    book = load_workbook(workbook)
    book.active["A4"] = formula
    book.save(workbook)
    with pytest.raises(OperationError, match="外部资源"):
        preview_number(workbook, NumberSettings())
