import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

from word_text_tool.models import ReplacementRule
from word_text_tool.replacement_engine import (
    RuleValidationError,
    find_matches,
    replace_runs,
    replace_text,
    validate_rule,
)


def test_replacement_is_non_cascading_and_not_duplicated():
    rules = (ReplacementRule("xxx公司", "yyy公司"), ReplacementRule("yyy公司", "zzz公司"))
    outcome = replace_text("甲方：xxx公司", rules)
    assert outcome.text == "甲方：yyy公司"
    assert outcome.count == 1


def test_leftmost_match_uses_longest_rule_at_same_position():
    rules = (ReplacementRule("项目", "A"), ReplacementRule("项目编号", "B"))
    assert replace_text("项目编号：01", rules).text == "B：01"


def test_same_length_rules_use_original_rule_order():
    rules = (ReplacementRule("编号", "first"), ReplacementRule("编号", "second"))
    matches = find_matches("编号", rules)
    assert matches[0].replacement == "first"


def test_match_advances_past_entire_range_and_skips_overlap():
    rule = ReplacementRule("aba", "X")
    matches = find_matches("ababa", (rule,))
    assert len(matches) == 1
    assert (matches[0].start, matches[0].end) == (0, 3)
    assert replace_text("ababa", (rule,)).text == "Xba"


def test_soft_line_break_is_matched_exactly():
    rule = ReplacementRule("第一行\n第二行", "合并")
    assert replace_text("第一行\n第二行", (rule,)).text == "合并"


def test_empty_replacement_deletes_match():
    assert replace_text("甲删除乙", (ReplacementRule("删除", ""),)).text == "甲乙"


@pytest.mark.parametrize("before", ["", "   ", "\t\n"])
def test_rule_rejects_empty_or_whitespace_only_before(before):
    with pytest.raises(RuleValidationError, match="不能为空或纯空白"):
        validate_rule(before, "新值", ())


def test_rule_preserves_significant_surrounding_whitespace():
    rule = validate_rule(" 项目编号 ", " 新编号 ", ())
    assert rule.before == " 项目编号 "
    assert rule.after == " 新编号 "


def test_rule_rejects_duplicate_before_text():
    existing = (ReplacementRule("项目编号", "A"),)
    with pytest.raises(RuleValidationError, match="已存在"):
        validate_rule("项目编号", "B", existing)


def test_rule_rejects_text_longer_than_10000_characters():
    with pytest.raises(RuleValidationError, match="10000"):
        validate_rule("甲" * 10001, "乙", ())


def test_rule_accepts_10000_characters_on_both_sides():
    rule = validate_rule("甲" * 10000, "乙" * 10000, ())
    assert len(rule.before) == len(rule.after) == 10000


def test_rule_rejects_replacement_longer_than_10000_characters():
    with pytest.raises(RuleValidationError, match="10000"):
        validate_rule("甲", "乙" * 10001, ())


def test_rule_matches_surrounding_spaces_and_punctuation_exactly():
    rule = ReplacementRule(" 项目! ", "替")
    assert replace_text(" 项目! 项目!", (rule,)).text == "替项目!"


def test_replace_across_runs_uses_first_run_format_without_duplication():
    paragraph = Document().add_paragraph()
    first = paragraph.add_run("xxx")
    first.font.name = "Microsoft YaHei UI"
    first.font.size = Pt(12)
    first.font.color.rgb = RGBColor(255, 0, 0)
    second = paragraph.add_run("公司及附件")
    second.bold = True

    count = replace_runs((first, second), (ReplacementRule("xxx公司", "yyy公司"),))

    assert count == 1
    assert paragraph.text == "yyy公司及附件"
    assert first.text == "yyy公司"
    assert first.font.name == "Microsoft YaHei UI"
    assert first.font.size == Pt(12)
    assert first.font.color.rgb == RGBColor(255, 0, 0)
    assert second.text == "及附件"
    assert second.bold is True


def test_replace_runs_handles_multiple_matches_in_one_run_right_to_left():
    paragraph = Document().add_paragraph()
    run = paragraph.add_run("甲旧值乙旧值丙")

    count = replace_runs((run,), (ReplacementRule("旧值", "新"),))

    assert count == 2
    assert run.text == "甲新乙新丙"


def test_replace_runs_handles_multiple_matches_across_distinct_runs():
    paragraph = Document().add_paragraph()
    first = paragraph.add_run("甲旧")
    second = paragraph.add_run("值乙")
    third = paragraph.add_run("旧值丙")

    count = replace_runs((first, second, third), (ReplacementRule("旧值", "新内容"),))

    assert count == 2
    assert paragraph.text == "甲新内容乙新内容丙"
    assert first.text == "甲新内容"
    assert second.text == "乙"
    assert third.text == "新内容丙"


def test_replace_runs_spans_four_runs_and_keeps_unaffected_run_format():
    paragraph = Document().add_paragraph()
    prefix = paragraph.add_run("前A")
    middle_one = paragraph.add_run("B")
    middle_two = paragraph.add_run("C")
    suffix = paragraph.add_run("D后")
    suffix.italic = True

    count = replace_runs(
        (prefix, middle_one, middle_two, suffix),
        (ReplacementRule("ABCD", "X"),),
    )

    assert count == 1
    assert paragraph.text == "前X后"
    assert [run.text for run in paragraph.runs] == ["前X", "", "", "后"]
    assert suffix.italic is True


def test_replace_runs_handles_empty_run_and_exact_run_boundaries():
    paragraph = Document().add_paragraph()
    first = paragraph.add_run("前")
    empty = paragraph.add_run("")
    second = paragraph.add_run("旧")
    third = paragraph.add_run("值")
    fourth = paragraph.add_run("后")

    count = replace_runs(
        (first, empty, second, third, fourth),
        (ReplacementRule("旧值", "新"),),
    )

    assert count == 1
    assert paragraph.text == "前新后"
    assert [run.text for run in paragraph.runs] == ["前", "", "新", "", "后"]


@pytest.mark.parametrize(
    ("replacement", "expected"),
    [("更长的新值", "甲更长的新值乙"), ("新", "甲新乙"), ("", "甲乙")],
)
def test_replace_runs_writes_length_changes_from_original_offsets(replacement, expected):
    paragraph = Document().add_paragraph()
    first = paragraph.add_run("甲旧")
    second = paragraph.add_run("值乙")

    count = replace_runs((first, second), (ReplacementRule("旧值", replacement),))

    assert count == 1
    assert paragraph.text == expected
    assert second.text == "乙"


def test_replace_runs_preserves_soft_break_semantics():
    paragraph = Document().add_paragraph()
    run = paragraph.add_run("第一行")
    run.add_break(WD_BREAK.LINE)
    run.add_text("第二行")

    count = replace_runs(
        (run,),
        (ReplacementRule("第一行\n第二行", "新一行\n新二行"),),
    )

    assert count == 1
    assert run.text == "新一行\n新二行"
    assert paragraph.text == "新一行\n新二行"
